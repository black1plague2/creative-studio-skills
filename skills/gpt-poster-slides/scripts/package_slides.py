#!/usr/bin/env python3
"""Package reviewed GPT poster images. Does not generate or edit their artwork."""
from __future__ import annotations
import argparse,json,pathlib,re,tempfile,os
from PIL import Image
from _deck_tools import load_manifest,contain,ensure_outputs

def package(manifest,output_dir,formats=('pdf','pptx'),overwrite=False,allow_unreviewed=False,output_stem=None):
    data,slides=load_manifest(manifest,allow_unreviewed)
    formats=tuple(dict.fromkeys(formats))
    if not formats or any(f not in ('pdf','pptx') for f in formats):raise ValueError('formats must contain pdf and/or pptx')
    stem=output_stem or re.sub(r'[^a-zA-Z0-9_-]+','-',data['title']).strip('-').lower()
    if not stem or not re.fullmatch(r'[a-zA-Z0-9_-]+',stem):raise ValueError('output stem must be a safe non-empty basename')
    draft=any(not s['reviewed'] for s in slides)
    if draft:stem='DRAFT-'+stem
    out=pathlib.Path(output_dir).resolve();out.mkdir(parents=True,exist_ok=True)
    names=[stem+'.'+f for f in formats]+[stem+'-transcript.md',stem+'-packaging-report.json']
    ensure_outputs([out/n for n in names],overwrite)
    color=data.get('canvas_color','#FFFFFF')
    with tempfile.TemporaryDirectory(prefix='.poster-package-',dir=out) as temp:
        temp=pathlib.Path(temp);rgb=[]
        for i,s in enumerate(slides):
            path=temp/f'{i:03}.png'
            with Image.open(s['resolved_image']) as im:
                im=im.convert('RGBA');base=Image.new('RGBA',im.size,color);base.alpha_composite(im);base.convert('RGB').save(path)
            rgb.append(path)
        if 'pdf' in formats:
            from reportlab.pdfgen import canvas
            from reportlab.lib.colors import HexColor
            from reportlab.lib.utils import ImageReader
            c=canvas.Canvas(str(temp/(stem+'.pdf')),pagesize=(720,405),pageCompression=1)
            c.setTitle(data['title'])
            for s,path in zip(slides,rgb):
                c.setFillColor(HexColor(color));c.rect(0,0,720,405,fill=1,stroke=0)
                check=s['image_check'];x,y,w,h=contain(check['width'],check['height'],720,405)
                c.drawImage(ImageReader(str(path)),x,y,width=w,height=h);c.showPage()
            c.save()
        if 'pptx' in formats:
            from pptx import Presentation
            from pptx.util import Inches
            from pptx.dml.color import RGBColor
            p=Presentation();p.slide_width=Inches(16);p.slide_height=Inches(9);p.core_properties.title=data['title']
            for s,path in zip(slides,rgb):
                slide=p.slides.add_slide(p.slide_layouts[6]);fill=slide.background.fill;fill.solid();fill.fore_color.rgb=RGBColor.from_string(color[1:])
                check=s['image_check'];x,y,w,h=contain(check['width'],check['height'],p.slide_width,p.slide_height)
                slide.shapes.add_picture(str(path),int(round(x)),int(round(y)),width=int(round(w)),height=int(round(h)))
                copy='\n'.join(s.get('expected_text',[]))
                slide.notes_slide.notes_text_frame.text=copy
            p.save(str(temp/(stem+'.pptx')))
        transcript=f"# {data['title']}\n\n"+('DRAFT: contains unreviewed images.\n\n' if draft else '')
        for i,s in enumerate(slides,1):transcript+=f"## Slide {i}: {s['id']}\n\n"+'\n\n'.join(s.get('expected_text',[]))+'\n\n'
        (temp/(stem+'-transcript.md')).write_text(transcript,encoding='utf-8')
        report={'title':data['title'],'draft':draft,'aspect_ratio':'16:9','slide_count':len(slides),'formats':list(formats),'order':[s['id'] for s in slides],'images':[{'id':s['id'],**s['image_check'],'path':s['image']} for s in slides],'outputs':names,'scope':'Mechanical packaging. Exact-copy, source and visual checks are review records, not independently proven by this script.'}
        (temp/(stem+'-packaging-report.json')).write_text(json.dumps(report,indent=2),encoding='utf-8')
        for name in names:os.replace(temp/name,out/name)
    return report

def main():
    p=argparse.ArgumentParser(description='Package local reviewed poster images into an exact 16:9 PDF/PPTX.')
    p.add_argument('manifest');p.add_argument('--output-dir',required=True);p.add_argument('--formats',nargs='+',choices=['pdf','pptx'],default=['pdf','pptx']);p.add_argument('--output-stem');p.add_argument('--overwrite',action='store_true');p.add_argument('--allow-unreviewed',action='store_true');a=p.parse_args()
    try: result=package(a.manifest,a.output_dir,a.formats,a.overwrite,a.allow_unreviewed,a.output_stem)
    except (ValueError,OSError,ImportError,json.JSONDecodeError) as exc:p.exit(1,str(exc)+'\n')
    print(json.dumps(result,indent=2))

if __name__=='__main__':main()
